# -*- coding: utf-8 -*-
"""ساخت نسخهٔ متنی (report_source.md) و نسخهٔ قابل ویرایش (final_consultant_report.docx)
از همان منبعِ داده‌ای که PDF را می‌سازد (report_content.py).

اجرا:  python3 build_text_outputs.py
"""
import os
import sys

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
import report_content as C
import report_photos as P

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

HERE = os.path.dirname(os.path.abspath(__file__))
MD = os.path.join(HERE, "report_source.md")
DOCX = os.path.join(HERE, "final_consultant_report.docx")

NAVY = RGBColor(0x17, 0x32, 0x4D)
STEEL = RGBColor(0x42, 0x65, 0x7D)
DARK = RGBColor(0x22, 0x30, 0x3C)


# ================================================================== Markdown
def md_table(header, rows):
    out = ["| " + " | ".join(header) + " |",
           "| " + " | ".join(["---"] * len(header)) + " |"]
    for r in rows:
        out.append("| " + " | ".join(str(x).replace("|", "／") for x in r) + " |")
    return out


def fa(n):
    """ارقام فارسی برای هم‌خوانی با متن گزارش."""
    return str(n).translate(str.maketrans("0123456789", "۰۱۲۳۴۵۶۷۸۹"))


def status_cell(v):
    if v in C.STATUS_LABEL:
        return "{} ({})".format(C.STATUS_LABEL[v], C.STATUS_SHORT[v])
    return v


def build_md():
    L = []
    L.append("# " + C.TITLE + " — " + C.UNITS)
    L.append("")
    L.append("**" + C.COMPANY + "**  ")
    L.append("**مشاور:** " + C.CONSULTANT + "  ")
    L.append("**دوره گزارش:** " + C.PERIOD + "  ")
    L.append("**تاریخ تنظیم:** " + C.REPORT_DATE)
    L.append("")
    for i in range(1, 5):
        L.append("---")
        L.append("")
        L.append("## " + C.PAGE_TITLES[i])
        L.append("")
        if i == 1:
            L.append(C.INTRO)
            L.append("")
            L.append("### خلاصه بازدیدها و دستورهای ابلاغی دوره")
            L.append("")
            L += md_table(C.VISITS_HEADER, C.VISITS)
            L.append("")
            L.append("### شاخص‌ها")
            L.append("")
            for v, k in C.CARDS_P1:
                L.append("- **{}** — {}".format(v, k))
            L.append("")
        elif i == 2:
            L.append("### آمار روزانه پخت و توزیع غذا")
            L.append("")
            L += md_table(C.STATS_HEADER, C.STATS)
            L.append("")
            L.append("### سهمیه مصرف به ازای هر نفر")
            L.append("")
            L += md_table(C.STANDARDS_HEADER, C.STANDARDS)
            L.append("")
            L.append("### اقدامات کنترل مصرف و پیشگیری از هزینه اضافی")
            L.append("")
            L += md_table(C.COST_HEADER, C.COST)
            L.append("")
        elif i == 3:
            L.append("### اقدامات تعمیراتی و تجهیزاتی")
            L.append("")
            L += md_table(C.REPAIRS_HEADER,
                          [r[:3] + [status_cell(r[3])] for r in C.REPAIRS])
            L.append("")
            L.append("### موارد عدم انطباق ثبت‌شده در ارزیابی کنترل کیفیت")
            L.append("")
            L += md_table(C.NC_HEADER, [r[:3] + [status_cell(r[3])] for r in C.NC])
            L.append("")
            L.append("### خلاصه وضعیت اقدامات")
            L.append("")
            L += md_table(C.STATUS_SUMMARY_HEADER,
                          [[status_cell(r[0])] + r[1:] for r in C.STATUS_SUMMARY])
            L.append("")
            L.append("### اقدامات انجام‌شده در دوره")
            L.append("")
            for n, it in enumerate(C.ACHIEVEMENTS, 1):
                L.append("{}. {}".format(n, it))
            L.append("")
        else:
            L.append(C.SUMMARY_TEXT.replace("\n\n", "\n\n"))
            L.append("")
            L.append("### شاخص‌های کلیدی")
            L.append("")
            for v, k in C.CARDS_P4:
                L.append("- **{}** — {}".format(v, k))
            L.append("")
            L.append("### خلاصهٔ عملکرد به تفکیک حوزهٔ فعالیت")
            L.append("")
            L += md_table(C.PERF_HEADER, C.PERF)
            L.append("")
            L.append("### دستاوردهای شاخص دوره")
            L.append("")
            for it in C.ACHIEVEMENTS_SHORT:
                L.append("- " + it)
            L.append("")
            L.append("### اولویت‌های آتی")
            L.append("")
            for it in C.PRIORITIES:
                L.append("- " + it)
            L.append("")
            L.append("### جمع‌بندی نهایی")
            L.append("")
            L.append(C.CLOSING)
            L.append("")
    for idx in sorted(P.CATEGORIES):
        cat = P.CATEGORIES[idx]
        L.append("---")
        L.append("")
        L.append("## پیوست تصویری {} — {}".format(fa(idx), cat["title"]))
        L.append("")
        L.append(cat["desc"])
        L.append("")
        L.append("فایل‌های تصویر در `assets/photos/` نگهداری می‌شوند.")
        L.append("")
        L += md_table(["شماره", "شرح تصویر", "تاریخ ثبت", "فایل"],
                      [[fa(n), t, d, f] for n, (c, f, t, d)
                        in enumerate([(p[0], p[1], p[2], p[3]) for p in P.PHOTOS
                                      if p[0] == idx],
                                     start=1 + sum(len(P.by_category(k))
                                                   for k in range(1, idx)))])
        L.append("")
    open(MD, "w", encoding="utf-8").write("\n".join(L))
    print("نوشته شد: " + MD)


# ================================================================== DOCX
def set_rtl(par):
    pPr = par._p.get_or_add_pPr()
    bidi = OxmlElement("w:bidi")
    bidi.set(qn("w:val"), "1")
    pPr.append(bidi)
    par.alignment = WD_ALIGN_PARAGRAPH.RIGHT


def _run(par, text, size, bold=False, color=None):
    r = par.add_run(text)
    r.font.name = "Vazirmatn"
    r._element.rPr.rFonts.set(qn("w:cs"), "Vazirmatn")
    r._element.rPr.rFonts.set(qn("w:ascii"), "Vazirmatn")
    r._element.rPr.rFonts.set(qn("w:hAnsi"), "Vazirmatn")
    r.font.size = Pt(size)
    r.font.bold = bold
    if color is not None:
        r.font.color.rgb = color
    return r


def add_par(doc, text, size=10, bold=False, color=None, space_after=5):
    p = doc.add_paragraph()
    set_rtl(p)
    _run(p, text, size, bold, color)
    p.paragraph_format.space_after = Pt(space_after)
    return p


def add_table(doc, header, rows, size=8.5):
    t = doc.add_table(rows=1, cols=len(header))
    t.style = "Table Grid"
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    for i, h in enumerate(header):
        cell = t.rows[0].cells[i]
        cell.text = ""
        p = cell.paragraphs[0]
        set_rtl(p)
        _run(p, h, size, True, RGBColor(0xFF, 0xFF, 0xFF))
    for row in rows:
        cells = t.add_row().cells
        for i, v in enumerate(row):
            cells[i].text = ""
            p = cells[i].paragraphs[0]
            set_rtl(p)
            _run(p, str(v), size, False, DARK)
    doc.add_paragraph()
    return t


def build_docx():
    doc = Document()
    sec = doc.sections[0]
    sec.page_width, sec.page_height = Cm(21), Cm(29.7)
    sec.left_margin = sec.right_margin = Cm(2)
    sec.top_margin = sec.bottom_margin = Cm(2)
    st = doc.styles["Normal"]
    st.font.name = "Vazirmatn"
    st.font.size = Pt(10)
    st.element.rPr.rFonts.set(qn("w:cs"), "Vazirmatn")

    add_par(doc, C.TITLE + " — " + C.UNITS, 15, True, NAVY)
    add_par(doc, C.COMPANY, 10.5, True, STEEL)
    add_par(doc, "مشاور: " + C.CONSULTANT + "    |    دوره گزارش: " + C.PERIOD
            + "    |    تاریخ تنظیم: " + C.REPORT_DATE, 9, False, STEEL)

    add_par(doc, C.PAGE_TITLES[1], 12.5, True, NAVY)
    add_par(doc, C.INTRO, 10)
    add_par(doc, "خلاصه بازدیدها و دستورهای ابلاغی دوره", 11, True, STEEL)
    add_table(doc, C.VISITS_HEADER, C.VISITS, 8)
    add_par(doc, "شاخص‌ها: " + " | ".join(v + " " + k for v, k in C.CARDS_P1),
            9.5, True, NAVY)

    add_par(doc, C.PAGE_TITLES[2], 12.5, True, NAVY)
    add_par(doc, "آمار روزانه پخت و توزیع غذا", 11, True, STEEL)
    add_table(doc, C.STATS_HEADER, C.STATS, 8)
    add_par(doc, "سهمیه مصرف به ازای هر نفر", 11, True, STEEL)
    add_table(doc, C.STANDARDS_HEADER, C.STANDARDS, 8)
    add_par(doc, "اقدامات کنترل مصرف و پیشگیری از هزینه اضافی", 11, True, STEEL)
    add_table(doc, C.COST_HEADER, C.COST, 8)

    add_par(doc, C.PAGE_TITLES[3], 12.5, True, NAVY)
    add_par(doc, "اقدامات تعمیراتی و تجهیزاتی", 11, True, STEEL)
    add_table(doc, C.REPAIRS_HEADER,
              [r[:3] + [C.STATUS_LABEL.get(r[3], r[3])] for r in C.REPAIRS], 8.5)
    add_par(doc, "موارد عدم انطباق ثبت‌شده در ارزیابی کنترل کیفیت", 11, True, STEEL)
    add_table(doc, C.NC_HEADER,
              [r[:3] + [C.STATUS_LABEL.get(r[3], r[3])] for r in C.NC], 8.5)
    add_par(doc, "خلاصه وضعیت اقدامات", 11, True, STEEL)
    add_table(doc, C.STATUS_SUMMARY_HEADER,
              [[C.STATUS_LABEL.get(r[0], r[0])] + r[1:] for r in C.STATUS_SUMMARY], 8.5)
    add_par(doc, "اقدامات انجام‌شده در دوره", 11, True, STEEL)
    for it in C.ACHIEVEMENTS:
        p = doc.add_paragraph(it, style="List Number")
        set_rtl(p)
        for r in p.runs:
            r.font.name = "Vazirmatn"
            r.font.size = Pt(9)
            r._element.rPr.rFonts.set(qn("w:cs"), "Vazirmatn")
    add_par(doc, C.PAGE_TITLES[4], 12.5, True, NAVY)
    for blk in C.SUMMARY_TEXT.split("\n\n"):
        add_par(doc, blk, 10)
    add_par(doc, "شاخص‌های کلیدی", 11, True, STEEL)
    add_table(doc, ["شاخص", "مقدار"],
              [[k, v] for v, k in C.CARDS_P4], 9)
    add_par(doc, "خلاصهٔ عملکرد به تفکیک حوزهٔ فعالیت", 11, True, STEEL)
    add_table(doc, C.PERF_HEADER, C.PERF, 8.5)
    add_par(doc, "دستاوردهای شاخص دوره", 11, True, STEEL)
    for it in C.ACHIEVEMENTS_SHORT:
        p = doc.add_paragraph(it, style="List Bullet")
        set_rtl(p)
        for r in p.runs:
            r.font.name = "Vazirmatn"
            r.font.size = Pt(9.5)
            r._element.rPr.rFonts.set(qn("w:cs"), "Vazirmatn")
    add_par(doc, "اولویت‌های آتی", 11, True, STEEL)
    for it in C.PRIORITIES:
        p = doc.add_paragraph(it, style="List Bullet")
        set_rtl(p)
        for r in p.runs:
            r.font.name = "Vazirmatn"
            r.font.size = Pt(9.5)
            r._element.rPr.rFonts.set(qn("w:cs"), "Vazirmatn")
    add_par(doc, "جمع‌بندی نهایی", 11, True, STEEL)
    add_par(doc, C.CLOSING, 10)

    # ---------------- پیوست تصویری (فهرست؛ تصاویر در PDF جای‌گذاری شده‌اند)
    for idx in sorted(P.CATEGORIES):
        cat = P.CATEGORIES[idx]
        add_par(doc, "پیوست تصویری {} — {}".format(fa(idx), cat["title"]), 12.5, True, NAVY)
        add_par(doc, cat["desc"], 9)
        add_par(doc, "فایل‌های تصویر در پوشهٔ assets/photos نگهداری می‌شوند.", 8.5,
                False, STEEL)
        first = 1 + sum(len(P.by_category(k)) for k in range(1, idx))
        rows = [[fa(first + i), it[2], it[3], it[1]]
                for i, it in enumerate(P.by_category(idx))]
        add_table(doc, ["شماره", "شرح تصویر", "تاریخ ثبت", "فایل"], rows, 8)

    doc.save(DOCX)
    print("نوشته شد: " + DOCX)


if __name__ == "__main__":
    build_md()
    build_docx()
